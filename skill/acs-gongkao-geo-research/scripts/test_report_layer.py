#!/usr/bin/env python3
"""Report Layer and Chinese customer-report focused regression.

Covers:
1. Empty Report Layer must fail independently of Measurement.
2. Complete Report Layer must pass strict validation.
3. A heading-only customer section must fail.
4. Developer/internal fields injected into the customer DOCX must fail.
"""
from __future__ import annotations
import importlib.util,json,sys,tempfile
from pathlib import Path
HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))


def load(name,file):
    spec=importlib.util.spec_from_file_location(name,HERE/file);m=importlib.util.module_from_spec(spec);sys.modules[name]=m;spec.loader.exec_module(m);return m


core=load('core_for_report_layer','test_v22_core.py')
val=load('val_for_report_layer','validate_run.py')
modeler=load('model_for_report_layer','build_report_model.py')

REQUIRED_CODES={"empty-asset-readiness","empty-concept-ownership","empty-analysis","empty-report-section","missing-report-docx"}
REPORT_LAYER_FILES=("asset_inputs.csv","asset_scores.csv","concept_ownership.csv","report_research_manifest.csv","report_layer_adjustments.json","analysis.json")


def errors(issues):return [x for x in issues if x.level=="error"]
def blocking(issues):return [x for x in issues if x.level in {"error","warning"}]


def strip_section(docx_in:Path,docx_out:Path,heading:str):
    from docx import Document
    from docx.oxml.ns import qn
    d=Document(str(docx_in));body=d.element.body;drop=False;removed=0
    for el in list(body):
        if el.tag==qn('w:p'):
            ppr=el.find(qn('w:pPr'));style=""
            if ppr is not None:
                st=ppr.find(qn('w:pStyle'))
                if st is not None:style=st.get(qn('w:val')) or ""
            text="".join(t.text or "" for t in el.iter(qn('w:t'))).strip()
            if style=="Heading1" and text==heading:drop=True;continue
            if drop and style=="Heading1":drop=False;continue
        if drop and el.tag in (qn('w:p'),qn('w:tbl')):body.remove(el);removed+=1
    if not removed:raise AssertionError(f"未能定位/清空章节正文：{heading}")
    d.save(str(docx_out))


def inject_internal_field(docx_in:Path,docx_out:Path):
    from docx import Document
    d=Document(str(docx_in));d.add_paragraph("measurement_target=institution");d.save(str(docx_out))


def main():
    checks=0;skips=[]
    with tempfile.TemporaryDirectory() as td:
        r=Path(td);core.fixture(r)
        codes={x.code for x in blocking(val.validate(r,True,'report'))}
        assert codes<={"missing-report-docx","missing-report-chart"}, f"完整 Report Layer fixture 数据层仍有阻塞项：{sorted(codes-{'missing-report-docx','missing-report-chart'})}"
        assert "missing-report-docx" in codes;checks+=1
        for n in REPORT_LAYER_FILES:(r/n).unlink(missing_ok=True)
        modeler.build(r)
        issues=val.validate(r,stage='report');errs=errors(issues);got={x.code for x in errs}
        assert errs,"空壳 Report Layer 未产生任何 error"
        missing=REQUIRED_CODES-got;assert not missing,f"空壳 Report Layer 未命中强制错误码 {sorted(missing)}；实际={sorted(got)}";checks+=1
        assert (1 if blocking(issues) else 0)!=0,"空壳 Report Layer 的 strict 退出码必须非 0";checks+=1
        codes2={x.code for x in blocking(val.validate(r,True,'report'))}
        assert "empty-report-section" in codes2 and "empty-analysis" in codes2,f"空壳 Report Layer 的 strict 仍不阻断：{sorted(codes2)}";checks+=1
        assert not blocking(val.validate(r,True,'measurement')),"负向 fixture 的 Measurement 层被破坏，无法证明 Report Gate 独立生效";checks+=1

    with tempfile.TemporaryDirectory() as td:
        r=Path(td);core.fixture(r)
        try:
            import matplotlib  # noqa: F401
            import docx as _docx  # noqa: F401
        except ImportError as e:
            skips.append(f'Report Layer DOCX 集成测试 SKIP: {type(e).__name__}: {e}')
        else:
            charts=load('charts_for_report_layer','generate_charts.py');docxmod=load('docx_for_report_layer','generate_report_docx.py')
            chart_result=charts.generate(r);assert len(chart_result.get('files') or [])==6,chart_result
            modeler.build(r);docxmod.render(r,r/'deliverables'/'report.docx')
            pos=blocking(val.validate(r,True,'report'))
            assert not pos,f"完整中文客户报告未通过 report strict：{[(x.code,x.message) for x in pos][:8]}";checks+=1

            full=r/'deliverables'/'report.docx';broken=r/'deliverables'/'report_empty_section.docx'
            strip_section(full,broken,"十一、如何理解本报告")
            full.rename(r/'deliverables'/'report_full.docx');broken.rename(full)
            codes3={x.code for x in blocking(val.validate(r,True,'report'))}
            assert "empty-report-section" in codes3,f"DOCX 空壳章节未被阻断：{sorted(codes3)}";checks+=1
            full.unlink();(r/'deliverables'/'report_full.docx').rename(full)

            leaked=r/'deliverables'/'report_internal_leak.docx';inject_internal_field(full,leaked)
            full.rename(r/'deliverables'/'report_clean.docx');leaked.rename(full)
            codes4={x.code for x in blocking(val.validate(r,True,'report'))}
            assert "customer-report-internal-field" in codes4 or "customer-report-snake-case" in codes4,f"客户报告内部字段泄漏未被阻断：{sorted(codes4)}";checks+=1
            full.unlink();(r/'deliverables'/'report_clean.docx').rename(full)

    print(f'PASS: GEO v2.2.1 Report Layer / 中文客户报告回归 {checks} 项通过')
    for s in skips:print(s)
    return 0


if __name__=='__main__':raise SystemExit(main())
