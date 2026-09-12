#!/usr/bin/env python3
"""Final DOCX QA for the Chinese customer report.

Checks both structural integrity and product readability: Chinese-first headings, no developer
fields/files/raw objects, citation N/A semantics, real charts/tables, and report-model consistency.
"""
from __future__ import annotations
import argparse,json,sys,zipfile
from collections import Counter
from pathlib import Path
HERE=Path(__file__).resolve().parent
sys.path.insert(0,str(HERE))


def main():
    ap=argparse.ArgumentParser();ap.add_argument("run_dir",type=Path);a=ap.parse_args();run=a.run_dir
    from docx import Document
    import importlib.util
    spec=importlib.util.spec_from_file_location("vrc",HERE/"validation_report_customer.py");vrc=importlib.util.module_from_spec(spec)
    sys.modules["vrc"]=vrc;spec.loader.exec_module(vrc)
    docx=run/"deliverables"/"report.docx";model=json.loads((run/"report_model.json").read_text(encoding="utf-8"))
    report={"file":str(docx),"size":docx.stat().st_size,"checks":[],"sections":[],"warnings":[]}
    def chk(name,ok,detail=""):
        report["checks"].append({"check":name,"ok":bool(ok),"detail":detail});return ok

    d=Document(str(docx));chk("DOCX 可正常打开",True,f"paragraphs={len(d.paragraphs)}, tables={len(d.tables)}")
    with zipfile.ZipFile(docx) as z:
        bad=z.testzip();chk("ZIP 结构无损坏",bad is None,str(bad))
        names=z.namelist();imgs=[n for n in names if n.startswith("word/media/")]
        chk("客户报告内嵌 6 张决策图",len(imgs)==6,f"{len(imgs)} 张：{[Path(x).name for x in imgs]}")
        chk("页眉页脚已写入",any(n.startswith("word/header") for n in names) and any(n.startswith("word/footer") for n in names))
        xml=z.read("word/document.xml").decode("utf-8",errors="ignore")

    ph=[x for x in ("待归纳","TODO","TBD","placeholder","占位文案","待补充","lorem") if x in xml]
    chk("无占位/调试文案",not ph,",".join(ph) or "none")

    sections=vrc.core.docx_sections(docx);rep={s["heading"]:s["content"] for s in sections}
    for heading in vrc.REQUIRED_DOCX_SECTIONS:
        ok=heading in rep and rep[heading]>0
        report["sections"].append({"heading":heading,"content_elements":rep.get(heading,0),"present":heading in rep})
        chk(f"章节「{heading}」存在且非空",ok,f"content={rep.get(heading,0)}")
    dup=[k for k,v in Counter(s["heading"] for s in sections).items() if v>1]
    chk("无重复一级章节",not dup,",".join(dup) or "none")

    empty_tables=0;tbl_rows=[]
    for t in d.tables:
        tbl_rows.append(len(t.rows))
        if len(t.rows)<2:empty_tables+=1
    chk("无只有表头的数据表",empty_tables==0,f"empty={empty_tables} / tables={len(d.tables)}")
    chk("客户报告数据表数量合理",len(d.tables)>=8 and sum(tbl_rows)>=30,f"tables={len(d.tables)}, rows_total={sum(tbl_rows)}")

    text="\n".join(p.text for p in d.paragraphs)
    for t in d.tables:
        for r in t.rows:text+="\n"+" | ".join(c.text for c in r.cells)

    customer_issues=[];vrc.check_customer_docx(run,customer_issues)
    chk("客户版中文化 / 无工程字段泄漏",not customer_issues,"; ".join(f"{x.code}:{x.message}" for x in customer_issues) or "PASS")
    chk("不展示工程审计章节","Research Audit" not in text and "报告层数据来源" not in text and "研究资产清单" not in text)
    chk("重点主体诊断包含优先动作","优先动作" in text)

    unscored=[]
    for x in model.get("asset_readiness") or []:
        try:float(x.get("asset_readiness"))
        except:unscored.append(x)
    if unscored:
        chk("证据不足主体明确写暂不评分","暂不评分" in text,f"unscored={len(unscored)}")
    else:
        chk("证据不足主体明确写暂不评分",True,"本轮所有报告主体均有可评分资产证据，无需出现“暂不评分”")

    answers=[];apath=run/"ai_answers.jsonl"
    if apath.is_file():
        for line in apath.read_text(encoding="utf-8").splitlines():
            if line.strip():answers.append(json.loads(line))
    total_citations=sum(len(a.get("citations") or []) for a in answers)
    if total_citations==0:chk("无引用链时显示不适用","不适用" in text)

    national=(model.get("ai_visibility") or {}).get("national_benchmarks") or []
    if national:
        r=national[0];a=r.get("mentioned_answers");b=r.get("answer_cells")
        if a not in (None,"") and b not in (None,""):chk("主榜命中/样本与模型对账",f"{a}/{b}" in text,f"value={a}/{b}")
    rb=model.get("robustness") or {}
    try:
        v=float(rb.get("positive_persistence_3of3_rate"));formatted=f"{(v*100 if abs(v)<=1 else v):.1f}%"
        chk("稳定性关键数字进入正文",formatted in text,f"value={formatted}")
    except Exception:pass

    fails=[c for c in report["checks"] if not c["ok"]];report["verdict"]="PASS" if not fails else "FAIL";report["failed"]=fails
    evidence_dir=run.parent.parent/"evidence";evidence_dir.mkdir(parents=True,exist_ok=True)
    (evidence_dir/"docx_qa.json").write_text(json.dumps(report,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"verdict":report["verdict"],"checks":len(report["checks"]),"failed":len(fails),"sections":len(report["sections"]),"images":len(imgs),"tables":len(d.tables)},ensure_ascii=False))
    for f in fails:print("  FAIL:",f["check"],"|",f["detail"])
    return 0 if not fails else 1

if __name__=="__main__":raise SystemExit(main())
