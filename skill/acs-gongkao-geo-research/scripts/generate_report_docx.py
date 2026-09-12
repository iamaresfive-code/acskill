#!/usr/bin/env python3
"""Single official renderer for GEO v2.2: stable A4 DOCX consulting report.

v2.2 Report Layer Completion：结构对齐 references/report-template.md（16 节），
且**拒绝输出空壳报告**——关键章节为空时直接抛错，不再用占位文案兜底。
"""
from __future__ import annotations
import argparse,json
from pathlib import Path
from fs_utils import ensure_directory

REQUIRED_NONEMPTY_LIST=("executive_summary","diagnoses","strategy","plan_90_days","risks")
REQUIRED_NONEMPTY_MAP=("measurement_protocol","robustness")


def _deps():
    try:
        from docx import Document
        from docx.shared import Mm,Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        return Document,Mm,Pt,WD_ALIGN_PARAGRAPH,WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT,OxmlElement,qn
    except Exception as e:raise RuntimeError(f"需要 python-docx：{e}")


def _assert_complete(model:dict):
    for key in REQUIRED_NONEMPTY_LIST:
        if not (model.get(key) or []):raise ValueError(f"report_model.{key} 为空，拒绝生成空壳报告")
    for key in REQUIRED_NONEMPTY_MAP:
        if not (model.get(key) or {}):raise ValueError(f"report_model.{key} 为空，拒绝生成空壳报告")
    if not (model.get("asset_readiness") or []):raise ValueError("report_model.asset_readiness 为空，拒绝生成空壳报告")
    if not (model.get("concept_map") or []):raise ValueError("report_model.concept_map 为空，拒绝生成空壳报告")
    if not (model.get("charts") or {}):raise ValueError("report_model.charts 为空，拒绝生成空壳报告")


def render(run:Path,out:Path):
    Document,Mm,Pt,ALIGN,TALIGN,CVAL,OxmlElement,qn=_deps();model=json.loads((run/"report_model.json").read_text(encoding="utf-8"))
    _assert_complete(model)
    meta=model.get("meta",{});d=Document();sec=d.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297);sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=Mm(19);sec.right_margin=Mm(19)
    styles=d.styles
    for sname,size,bold in [("Normal",10.5,False),("Title",30,True),("Heading 1",20,True),("Heading 2",15,True),("Heading 3",12,True)]:
        s=styles[sname];s.font.name="Aptos";s.font.size=Pt(size);s.font.bold=bold;s._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
    styles["Normal"].paragraph_format.line_spacing=1.45;styles["Normal"].paragraph_format.space_after=Pt(5)
    def set_run(r,size=None,bold=None):
        r.font.name="Aptos";r._element.rPr.rFonts.set(qn("w:eastAsia"),"Microsoft YaHei")
        if size:r.font.size=Pt(size)
        if bold is not None:r.bold=bold
    def page_break():d.add_page_break()
    def h(text,level=1):
        p=d.add_paragraph(style=f"Heading {level}");p.paragraph_format.keep_with_next=True;p.add_run(text);return p
    def para(text,size=10.5,italic=False):
        p=d.add_paragraph();r=p.add_run(str(text));set_run(r,size);r.italic=italic;return p
    def bullet(text):
        p=d.add_paragraph(style="List Bullet");p.add_run(str(text));return p
    def set_repeat_header(row):
        trPr=row._tr.get_or_add_trPr();tblHeader=OxmlElement("w:tblHeader");tblHeader.set(qn("w:val"),"true");trPr.append(tblHeader)
    def no_split(row):
        trPr=row._tr.get_or_add_trPr();x=OxmlElement("w:cantSplit");trPr.append(x)
    def shade(cell,fill="EAF0F6"):
        tcPr=cell._tc.get_or_add_tcPr();shd=OxmlElement("w:shd");shd.set(qn("w:fill"),fill);tcPr.append(shd)
    def table(headers,rows,widths=None,font=8.5):
        t=d.add_table(rows=1,cols=len(headers));t.alignment=TALIGN.CENTER;t.style="Table Grid";set_repeat_header(t.rows[0])
        for i,x in enumerate(headers):
            c=t.rows[0].cells[i];c.text=str(x);shade(c);c.vertical_alignment=CVAL.CENTER
            for p in c.paragraphs:
                for r in p.runs:set_run(r,font,True)
        for row in rows:
            cells=t.add_row().cells;no_split(t.rows[-1])
            for i,v in enumerate(row):
                cells[i].text=str(v if v is not None else "");cells[i].vertical_alignment=CVAL.TOP
                for p in cells[i].paragraphs:
                    for r in p.runs:set_run(r,font)
        if widths:
            for row in t.rows:
                for i,w in enumerate(widths):row.cells[i].width=Mm(w)
        return t
    def chart(key):
        rel=(model.get("charts") or {}).get(key);p=run/rel if rel else None
        if not p or not p.is_file():raise ValueError(f"图表缺失：{rel}")
        d.add_picture(str(p),width=Mm(165));d.paragraphs[-1].alignment=ALIGN.CENTER
    def pct(v):
        try:return f"{float(v)*100:.1f}%"
        except:return "—"
    def hit(r):
        m=r.get("target_metrics") or {};k=r.get("measurement_target_for_report") or ""
        row=m.get(k) or (next(iter(m.values())) if m else {})
        a=row.get("mentioned_answers");b=row.get("answer_cells");return f"{a}/{b}" if a not in (None,"") and b not in (None,"") else "—"
    def vis_rows(rows):
        return [[i+1,r.get("canonical_name"),hit(r),pct(r.get("nomination_rate")),pct(r.get("top3_rate")),pct(r.get("first_mention_rate")),pct(r.get("citation_rate")),r.get("asset_readiness") or "—"] for i,r in enumerate(rows)]
    VIS_HEAD=["排名","主体","提名(命中/样本)","提名率","Top3率","首提率","引用率","资产成熟度"]
    VIS_W=[11,32,22,18,18,18,18,25]

    for _ in range(5):d.add_paragraph("")
    p=d.add_paragraph();p.alignment=ALIGN.CENTER;r=p.add_run("INTERNAL RESEARCH");set_run(r,10,True)
    p=d.add_paragraph();p.alignment=ALIGN.CENTER;r=p.add_run(model.get("title",""));set_run(r,28,True)
    p=d.add_paragraph();p.alignment=ALIGN.CENTER;r=p.add_run(model.get("subtitle",""));set_run(r,13)
    for text in [f"研究范围：{meta.get('normalized_region') or meta.get('requested_region','')}",f"观察日期：{meta.get('observation_date','')}","正式交付：Word / DOCX","研究边界：GEO 不代表教学质量、市场份额或真实口碑"]:
        p=d.add_paragraph();p.alignment=ALIGN.CENTER;p.add_run(text)
    from docx.enum.section import WD_SECTION
    section=d.add_section(WD_SECTION.NEW_PAGE);section.page_width=Mm(210);section.page_height=Mm(297);section.top_margin=Mm(18);section.bottom_margin=Mm(18);section.left_margin=Mm(19);section.right_margin=Mm(19);section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
    hp=section.header.paragraphs[0];hp.text=model.get("title","");hp.alignment=ALIGN.CENTER
    fp=section.footer.paragraphs[0];fp.alignment=ALIGN.CENTER;fp.add_run("公开信息研究 · ");fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");fp._p.append(fld)
    sec.header.paragraphs[0].clear();sec.footer.paragraphs[0].clear()

    # 研究概览 / KPI
    h("研究概览",1);k=model.get("kpis",{});table(["正式研究主体","全国基准","本地/区域机构","Expert/IP","观察组","AI引擎"],[[k.get("included_entities"),k.get("national_benchmarks"),k.get("local_institutions"),k.get("expert_ip"),k.get("observation_entities"),k.get("ai_engines")]],widths=[28,28,30,26,26,26],font=9)
    para("本报告只回答两件事：在固定无品牌问题的真实 AI 回答中，谁被提到、以什么强度和顺序被提到；以及公开互联网资产为什么会让机器更容易认识、理解和引用它。GEO 不代表教学质量、通过率、招生量、市场份额或一般口碑。",9.5)

    # Executive Summary
    h("Executive Summary",1)
    for x in model.get("executive_summary",[]):bullet(x)

    # 一、研究范围与测量协议
    page_break();h("一、研究范围与测量协议",1)
    pr=model.get("measurement_protocol") or {}
    table(["协议字段","取值"],[["研究地区",meta.get("normalized_region") or meta.get("requested_region","")],["研究模式",meta.get("research_mode","")],["measurement_profile",pr.get("measurement_profile","")],["sampling_mode",pr.get("sampling_mode","")],["answer_context_mode",pr.get("answer_context_mode_expected","")],["context_isolation_level",pr.get("context_isolation_level","")],["query_variant_mode",pr.get("query_variant_mode","")],["repeat_runs_expected",pr.get("repeat_runs_expected","")],["fresh_context_required",pr.get("fresh_context_required","")],["page_collection_status",pr.get("page_collection_status","")],["observation_date",pr.get("observation_date","")],["sampled_at_range",pr.get("sampled_at_range","")],["answer_cells",pr.get("answer_cells","")],["AI 引擎",", ".join(pr.get("ai_engines_sampled") or [])]],widths=[52,108],font=9)
    h("强制方法披露",2)
    for x in (pr.get("declarations") or []):bullet(x)
    if pr.get("fresh_context_note"):para(f"fresh context 说明：{pr.get('fresh_context_note')}",9)
    if pr.get("ai_engine_notes"):para(f"引擎可用性说明：{pr.get('ai_engine_notes')}",9)
    h("数据口径",2)
    para("AI Answer Measurement 与 Open-Web SERP / 网页提及物理分离；SERP 只用于发现与解释，不替代真实 AI 提名。Metric 一行 = 主体 × measurement_target；institution 与 IP 分母不混合。Asset Readiness 与 Concept Ownership 只来自可核验公开证据，未取到证据的维度记为 unknown。",9.5)

    # 二、Market Universe
    page_break();h("二、Market Universe / 本次到底研究谁",1);chart("universe")
    table(["主体","类型","市场范围","市场角色","用户Seed","纳入依据"],[[r.get("canonical_name"),r.get("entity_type"),r.get("market_scope"),r.get("market_role"),r.get("user_seed"),r.get("salience_basis")] for group in model.get("market_universe",{}).values() for r in group],widths=[34,22,22,32,18,42],font=8)

    # 三、全国品牌
    page_break();h("三、全国品牌在本地区的 AI GEO 表现",1);chart("national_visibility")
    rows=model.get("ai_visibility",{}).get("national_benchmarks",[]);table(VIS_HEAD,vis_rows(rows),widths=VIS_W,font=8.5)
    para("样本量口径：institution 组 60 个 Answer Cell / 主体；提名率为正向提名（recommended + listed）占该组 Answer Cell 的比例。Citation Rate 分母为已被正向提名的 Answer Cell。单引擎条件下 Engine Coverage Rate 恒为 100%、Cross-model Consistency 记 N.A.，不得表述为跨模型共识。",8.5)

    # 四、本土 / 区域机构
    page_break();h("四、本土 / 区域机构 AI GEO",1);chart("local_visibility")
    rows=model.get("ai_visibility",{}).get("local_institutions",[]);table(VIS_HEAD,vis_rows(rows),widths=VIS_W,font=8.5)
    para("本土/区域机构按 Stage 1 冻结的 market_role（local-core / local-active）与 market_scope（local / regional）分桶，不因 measurement_target 或 entity_type 重新分桶。",8.5)

    # 五、Expert / IP
    page_break();h("五、Expert / IP GEO",1);chart("ip_visibility")
    rows=model.get("ai_visibility",{}).get("expert_ip",[]);table(VIS_HEAD,vis_rows(rows),widths=VIS_W,font=8.5)
    para("IP 组分母为 24 个 Answer Cell / 主体。Hybrid 主体若同时具备 institution 与 IP 入口，在报告中按两套 target-specific 指标分别呈现，不在本表中合并分母。",8.5)

    # 六、Robustness
    page_break();h("六、Query / Retrieval Robustness（查询/检索鲁棒性）",1)
    rb=model.get("robustness") or {}
    table(["指标","数值"],[["查询变体模式",rb.get("query_variant_mode","")],["变体证据状态",json.dumps(rb.get("variant_evidence_status_counts") or {},ensure_ascii=False)],["解释口径",rb.get("interpretation","")],["positive_persistence_3of3_rate",f"{rb.get('positive_persistence_3of3_rate','')}（{rb.get('positive_persistence_numerator','')}/{rb.get('positive_persistence_denominator','')}）"],["Pairwise Positive-set Jaccard（mean / median）",f"{rb.get('pairwise_positive_set_jaccard_mean','')} / {rb.get('pairwise_positive_set_jaccard_median','')}"],["Exact Positive-set Match Rate",rb.get("exact_positive_set_match_rate","")]],widths=[70,90],font=9)
    hp2=rb.get("hit_pattern_distribution") or {}
    if hp2:table(["Hit Pattern","entity×query 数"],[[x,hp2[x]] for x in ("0/3","1/3","2/3","3/3") if x in hp2],widths=[40,40],font=9)
    para("positive_persistence_3of3_rate 只表示“至少一次正向命中的 entity×query 中三个变体全部正向命中的比例”，**不是** Overall Repeat Stability（0/N 稳定负例不在其分母）。本轮 query_variant_mode=semantic-retrieval-variants，因此本组指标只能解释为 Query/Retrieval Robustness。",9)
    if rb.get("note"):para(str(rb.get("note")),9)

    # 七、Concept Ownership
    page_break();h("七、概念山头 / Concept Ownership",1);chart("concept_ownership")
    concepts=model.get("concept_map",[])
    table(["概念","主体","绑定类型","强度","来源类型","可核验依据"],[[x.get("concept"),x.get("canonical_name") or x.get("entity_name"),x.get("binding_type"),x.get("strength"),x.get("source_type"),x.get("basis")] for x in concepts],widths=[32,30,28,11,20,51],font=8)
    para("绑定类型区分品牌主动表达（owned-declaration）、公开内容高频绑定、第三方描述与单次偶发提及；单次第三方提及不得包装为强 Concept Ownership。全部绑定均附来源与观测日期，见附录。",8.5)

    # 八、Asset Readiness
    page_break();h("八、GEO Asset Readiness / 为什么 AI 可能认识它",1);chart("asset_readiness")
    assets=model.get("asset_readiness",[])
    table(["排名","主体","资产成熟度","Tier","证据数","独立域名","自有来源依赖","unknown 维度"],[[i+1,x.get("canonical_name") or x.get("entity_id"),x.get("asset_readiness"),x.get("asset_tier"),x.get("evidence_count"),x.get("independent_domains"),pct(x.get("owned_source_dependency")),(x.get("unknown_fields") or "").replace("|","、")] for i,x in enumerate(assets)],widths=[10,34,17,11,14,16,20,24],font=8)
    para("资产成熟度由 7 个维度加权（entity_clarity 25 / regional_semantic_density 20 / open_web_assets 15 / external_authority 15 / content_depth_freshness 10 / data_tool_assets 10 / platform_coverage 5）计算。**没有找到证据 ≠ 0 分**：未取到证据的维度记为 unknown，不计入分母，并在表中显式列出；unknown 占比过高时 Tier 记 U（证据不足），不得据此判断主体“没有资产”。",9)

    # 九、观察组与 emergent
    page_break();h("九、AI 可见的观察组与 AI-emergent 竞争主体",1)
    obs=model.get("ai_visible_observation",[])
    if obs:
        h("Observation 中被 AI 实际提名的主体",2)
        table(["主体","提名(命中/样本)","提名率","Top3率","市场范围","Stage 1 状态"],[[x.get("canonical_name"),hit(x),pct(x.get("nomination_rate")),pct(x.get("top3_rate")),x.get("market_scope"),x.get("universe_status")] for x in obs],widths=[40,24,22,22,28,34],font=8.5)
        para("这些主体未进入 Stage 1 主榜，但真实 AI Answer 已出现提名，应在最终解释中单独复核是否需要升级为正式竞争主体。",9)
    else:para("本轮 Observation 组未出现正向提名。",9.5)
    ai_new=model.get("ai_visible_emergent",[])
    if ai_new:
        h("AI 回答自然带出的新竞争主体",2)
        table(["主体","Target","提名(命中/样本)","提名率","Top3率","市场范围"],[[x.get("canonical_name"),x.get("measurement_target_for_report") or x.get("measurement_target"),hit(x),pct(x.get("nomination_rate")),pct(x.get("top3_rate")),x.get("market_scope")] for x in ai_new],widths=[38,20,24,22,22,44],font=8)
        para("这些主体不在 Stage 1 已确认 Universe 中，而是在正式 AI Answer Measurement 中自然出现。它们不得回写污染预先确认的主榜，但必须保留、解析并单独披露。",9)
    else:para("本轮未出现 Universe 外已解析主体。",9.5)

    # 十、重点主体诊断
    page_break();h("十、重点主体诊断",1)
    for x in model["diagnoses"]:
        h(str(x.get("name") or x.get("entity") or "主体诊断"),2)
        table(["项目","判断"],[["市场角色",x.get("market_role","")],["AI表现",x.get("ai_visibility","")],["最强资产",x.get("strongest_asset","")],["主要短板",x.get("largest_gap","")],["关键动作",x.get("recommendation","")]],widths=[30,130],font=9)

    # 十一、策略与 90 天
    page_break();h("十一、区域策略与 90 天 GEO 工程",1)
    h("区域策略",2)
    for x in model["strategy"]:bullet(x)
    h("90 天工程",2)
    for x in model["plan_90_days"]:bullet(x)

    # 十二、风险与局限
    page_break();h("十二、风险与局限",1)
    for x in model["risks"]:bullet(x)

    # Appendix
    page_break();h("Appendix / Research Audit",1);app=model.get("appendix",{});rc=app.get("recheck",{});rr=app.get("resolution_recheck",{})
    table(["复判类型","记录数","分歧","已解决","Evidence"],[["Annotation Blind Recheck",rc.get("rows",0),rc.get("disagreements",0),rc.get("resolved_disagreements",0),k.get("evidence_count",0)],["Entity Resolution Recheck",rr.get("rows",0),"—",rr.get("confirmed",0),"—"]],widths=[52,24,20,22,30],font=9)
    h("Intent 分布（最终持久化数据实算）",2)
    table(["mention_intent","条数"],[[a,b] for a,b in sorted((model.get("measurement_qa",{}).get("annotation_intent_distribution") or {}).items())],widths=[40,40],font=9)
    h("研究方法边界",2)
    for x in app.get("methodology_notes",[]):bullet(x)
    h("报告层数据来源",2)
    for x in ["asset_inputs.csv（7 维资产评分 + 证据引用 + unknown 声明）","report_research_manifest.csv（本轮新增公开网页证据清单：来源分级、观测日期、HTTP 状态、正文哈希）","concept_ownership.csv（概念绑定：绑定类型 + 强度 + 来源 + 依据）","analysis.json（事实 / 指标 / 代理指标 / 推断 / 建议 五层分离）","score_assets.py → build_report_model.py → generate_report_docx.py（图表与统计均由脚本从最终持久化产物实算）"]:bullet(x)
    h("研究资产清单",2)
    for x in app.get("research_assets",[]):bullet(x)

    ensure_directory(out.parent);d.save(out);return out


def main():
    p=argparse.ArgumentParser();p.add_argument("run_dir",type=Path);p.add_argument("--output",type=Path);a=p.parse_args();out=a.output or a.run_dir/"deliverables"/"report.docx"
    try:print(render(a.run_dir,out));return 0
    except (OSError,RuntimeError,ValueError) as e:print(f"generate_report_docx：错误：{e}");return 2
if __name__=="__main__":raise SystemExit(main())
